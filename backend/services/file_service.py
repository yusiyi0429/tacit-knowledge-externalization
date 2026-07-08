"""File upload, caching, Excel and text extraction services."""

from __future__ import annotations

import json
import os
import uuid
from pathlib import Path

import openpyxl

from pipeline_artifacts import (
    basename_only,
    infer_file_step,
    is_step1_filename,
    is_step2_preextract_filename,
    is_step3_revision_filename,
    is_step3_final_filename,
    locate_workspace_file,
    resolve_client_excel_path,
    safe_workspace_path,
    workspace_path_for,
)
from shared import _safe_workbook, extract_text_from_file, extract_text_from_path


def save_upload_service(workspace: Path, file_obj, prefix: str = "upload", pipeline_id: str = "") -> str:
    """Save an uploaded file to workspace, return the path."""
    ext = Path(file_obj.filename).suffix if file_obj.filename else ".xlsx"
    fname = f"{prefix}_{uuid.uuid4().hex[:8]}{ext}"
    step = infer_file_step(fname) or "uploads"
    fpath = workspace_path_for(workspace, pipeline_id or "", step, fname)
    fpath.parent.mkdir(parents=True, exist_ok=True)
    file_obj.save(str(fpath))
    return str(fpath)


def cache_upload(workspace: Path, file_obj, pipeline_id: str, step: str, pipelines_lock, load_pipelines, save_pipelines):
    """Cache uploaded source file in workspace, return cached filename."""
    prefix = f"cache_s{step or 'x'}"
    saved_path = save_upload_service(workspace, file_obj, prefix=prefix, pipeline_id=pipeline_id)
    base = os.path.basename(saved_path)
    resp = {"status": "ok", "file_name": base}

    if pipeline_id and step in {"2", "3"}:
        with pipelines_lock:
            pipelines = load_pipelines()
            for p in pipelines:
                if p["id"] == pipeline_id:
                    sd = p.setdefault("step_data", {})
                    sd[f"step{step}_cached_file"] = base
                    sd[f"step{step}_cached_name"] = file_obj.filename
                    save_pipelines(pipelines)
                    break
    return resp


def trim_excel_rows(rows, max_rows=200, max_cols=40):
    """裁剪尾部空行空列，避免 Luckysheet 渲染超大稀疏表卡顿。"""
    if not rows:
        return [[]]
    trimmed = []
    for row in rows[:max_rows]:
        trimmed.append([cell if cell is not None else "" for cell in (row or [])])
    last_r = len(trimmed) - 1
    while last_r > 0:
        if any(str(c or "").strip() for c in trimmed[last_r]):
            break
        last_r -= 1
    trimmed = trimmed[: last_r + 1]
    if not trimmed:
        return [[]]
    last_c = 0
    for row in trimmed:
        for i, c in enumerate(row):
            if str(c or "").strip():
                last_c = max(last_c, i)
    last_c = min(last_c, max_cols - 1)
    out = []
    for row in trimmed:
        r = list(row[: last_c + 1])
        while len(r) <= last_c:
            r.append("")
        out.append(r)
    return out or [[]]


def read_excel_for_editor(workspace: Path, input_path: str, pipeline_id: str, step: str,
                          pipelines_lock, load_pipelines, save_pipelines):
    """Read Excel file and return structured data for online editing."""
    wb = None
    try:
        wb = openpyxl.load_workbook(input_path, data_only=True)
        sheets = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            raw_rows = []
            for row in ws.iter_rows(values_only=True):
                raw_rows.append([cell if cell is not None else "" for cell in row])
            rows = trim_excel_rows(raw_rows)
            merges = []
            for merge_range in list(ws.merged_cells.ranges):
                merges.append({
                    "r": merge_range.min_row - 1,
                    "c": merge_range.min_col - 1,
                    "rs": merge_range.max_row - merge_range.min_row + 1,
                    "cs": merge_range.max_col - merge_range.min_col + 1,
                })
            sheets[sheet_name] = {"rows": rows, "merges": merges}

        if pipeline_id:
            with pipelines_lock:
                pipelines = load_pipelines()
                for p in pipelines:
                    if p["id"] == pipeline_id:
                        sd = p.setdefault("step_data", {})
                        base = basename_only(input_path)
                        sd[f"step{step}_excel_path"] = base
                        if str(step) == "1" and is_step1_filename(base):
                            sd["step1_output_file"] = base
                            sd["step1_download_url"] = "/downloads/" + base
                        save_pipelines(pipelines)
                        break

        return {"status": "ok", "sheets": sheets, "file_path": basename_only(input_path)}
    except Exception as e:
        return {"status": "error", "error": f"读取 Excel 失败: {str(e)}"}
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def save_excel_from_editor(workspace: Path, file_path: str, file_name: str, sheets: dict,
                           pipeline_id: str, step: str, pipelines_lock, load_pipelines, save_pipelines):
    """Save edited Excel data back to file and return download URL."""
    resolved = resolve_client_excel_path(workspace, file_path, file_name, pipeline_id=pipeline_id)
    if not resolved:
        return {"status": "error", "error": "源文件不存在或路径非法，请重新打开编辑"}
    file_path = str(resolved)

    if not sheets:
        return {"status": "error", "error": "无数据可保存"}

    wb = None
    try:
        wb = openpyxl.load_workbook(file_path)

        def _sheet_payload(raw):
            if isinstance(raw, list):
                return raw, []
            if isinstance(raw, dict):
                return raw.get("rows", []), raw.get("merges", [])
            return [], []

        for sheet_name, raw in sheets.items():
            rows, merges = _sheet_payload(raw)
            if sheet_name not in wb.sheetnames:
                wb.create_sheet(sheet_name)
            ws = wb[sheet_name]

            for merge_range in list(ws.merged_cells.ranges):
                ws.unmerge_cells(str(merge_range))

            if ws.max_row:
                ws.delete_rows(1, ws.max_row)

            for r_idx, row_data in enumerate(rows, start=1):
                for c_idx, value in enumerate(row_data, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=value if value != "" else None)

            for m in merges:
                rs = int(m.get("rs", 1))
                cs = int(m.get("cs", 1))
                if rs > 1 or cs > 1:
                    r0 = int(m.get("r", 0)) + 1
                    c0 = int(m.get("c", 0)) + 1
                    ws.merge_cells(
                        start_row=r0,
                        start_column=c0,
                        end_row=r0 + rs - 1,
                        end_column=c0 + cs - 1,
                    )

        save_name = f"edited_step{step}_{uuid.uuid4().hex[:8]}.xlsx"
        save_path_obj = workspace_path_for(workspace, pipeline_id, infer_file_step(save_name) or f"step{step}", save_name)
        save_path_obj.parent.mkdir(parents=True, exist_ok=True)
        save_path = str(save_path_obj)
        wb.save(save_path)
        wb.save(file_path)

        if pipeline_id:
            with pipelines_lock:
                pipelines = load_pipelines()
                for p in pipelines:
                    if p["id"] == pipeline_id:
                        sd = p.setdefault("step_data", {})
                        base = basename_only(file_path)
                        sd[f"step{step}_excel_path"] = base
                        if str(step) == "1" and is_step1_filename(base):
                            sd["step1_output_file"] = base
                            sd["step1_download_url"] = "/downloads/" + base
                        elif str(step) == "2" and is_step2_preextract_filename(base):
                            sd["step2_output_file"] = base
                            sd["step2_download_url"] = "/downloads/" + base
                            for stale in ("step2_preview_name", "step2_preview_url", "step2_download"):
                                sd.pop(stale, None)
                        elif str(step) == "3" and is_step3_revision_filename(base):
                            sd["step3_revision_file"] = base
                            sd["step3_download_url"] = "/downloads/" + save_name
                        elif str(step) == "3" and is_step3_final_filename(base):
                            sd["step3_final_file"] = base
                            sd["step3_final_download_url"] = "/downloads/" + save_name
                        elif str(step) == "3" and is_step3_revision_filename(base):
                            sd["step3_revision_file"] = base
                            sd["step3_download_url"] = "/downloads/" + save_name
                        save_pipelines(pipelines)
                        break

        return {
            "status": "ok",
            "download_url": "/downloads/" + save_name,
            "file_path": basename_only(file_path),
            "file_name": basename_only(file_path),
            "message": "保存成功"
        }
    except Exception as e:
        return {"status": "error", "error": f"保存 Excel 失败: {str(e)}"}
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
