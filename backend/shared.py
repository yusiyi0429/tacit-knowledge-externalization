"""
Shared utilities, configuration, and global state for the tacit knowledge app.

Extracted from app_server.py to reduce module size and improve testability.
All route handlers import from this module.
"""

from __future__ import annotations

import datetime
import json
import logging
import os
import re
import shutil
import tempfile
import threading
import uuid
from contextlib import contextmanager
from pathlib import Path
from typing import Any

import yaml
import openpyxl
from flask import Flask, request, jsonify, send_file, send_from_directory, Response

from llm_client import (
    API_TYPE_CCB, API_TYPE_OPENAI, LlmApiError,
    call_llm, call_llm_with_retry, extract_assistant_content,
    iter_llm_stream, normalize_llm_url,
)
from pipeline_artifacts import (
    basename_only, downstream_output_keys, keys_to_clear_from_step,
    is_download_allowed, is_step1_filename, is_step2_preextract_filename,
    is_step3_revision_filename, is_step3_final_filename,
    resolve_cache_file_path, resolve_client_excel_path, safe_workspace_path,
    validate_step_data_patch, resolve_knowledge_workbook_path,
    PROTECTED_WORKSPACE_FILES,
)
from release_info import STEP2_EXCEL_BUILD, get_release_info
from knowledge_fusion import (
    merge_extraction_results, fuse_sources, detect_duplicates,
    detect_conflicts, interview_answers_to_records, aggregate_signals,
)
from interview_session import (
    execute_interview_session, collect_interview_answers,
    build_interview_prompt, INTERVIEW_METHODS,
)
from i18n import resolve_locale, t

# ─── Path Resolution ──────────────────────────────────────────────
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
SAMPLES_DIR = PROJECT_DIR / "data" / "samples"
CONFIG_DIR = PROJECT_DIR / "config"
FRONTEND_DIR = PROJECT_DIR / "frontend"
SCHEMA_PATH = CONFIG_DIR / "scenario-schema.yaml"
LLM_CONFIG_PATH = CONFIG_DIR / "llm-config.yaml"
LLM_CONFIG_LOCAL_PATH = CONFIG_DIR / "llm-config.local.yaml"

# Optional API token auth
APP_AUTH_TOKEN = os.environ.get("APP_AUTH_TOKEN", "").strip()

# Workspace
WORKSPACE = Path(os.environ.get("WORKSPACE_DIR", Path(tempfile.gettempdir()) / "tacit_knowledge_app"))
WORKSPACE.mkdir(parents=True, exist_ok=True)

# Persistence paths
CUSTOM_MODELS_PATH = WORKSPACE / "custom_models.json"
PRESET_OVERRIDES_PATH = WORKSPACE / "preset_overrides.json"
PIPELINES_PATH = WORKSPACE / "pipelines.json"

# Thread-safe state
_models_lock = threading.Lock()
_pipelines_lock = threading.Lock()

AUTH_EXEMPT_PATHS = frozenset({"/api/health", "/api/version", "/api/auth/config"})

# ─── Style Rules ───────────────────────────────────────────────────
EXTRACT_STYLE_RULES = {
    "标准萃取": {
        "temperature": 0.25, "max_tokens": 4096,
        "min_items": 8, "max_items": 22,
        "prompt_hint": "平衡覆盖核心规则、流程与经验，优先可执行条目。",
    },
    "深度萃取": {
        "temperature": 0.35, "max_tokens": 4096,
        "min_items": 12, "max_items": 40,
        "prompt_hint": "优先完整覆盖，尽量补全触发条件、判断逻辑、反模式与证据字段。",
    },
    "精简萃取": {
        "temperature": 0.15, "max_tokens": 2048,
        "min_items": 5, "max_items": 10,
        "prompt_hint": "只保留高价值高置信条目，减少冗余与重复。",
    },
}

REVISION_STYLE_RULES = {
    "标准修订": {
        "temperature": 0.25, "max_actions": 40,
        "allowed_actions": {"modify", "supplement", "add", "delete"},
        "prompt_hint": "平衡修订：采纳明确建议，同时保持原有合理内容。",
    },
    "严格修订": {
        "temperature": 0.1, "max_actions": 20,
        "allowed_actions": {"modify", "supplement"},
        "prompt_hint": "保守修订：仅处理证据充分、定位明确的修改/补充，禁止新增和删除。",
    },
    "宽松修订": {
        "temperature": 0.35, "max_actions": 80,
        "allowed_actions": {"modify", "supplement", "add", "delete"},
        "prompt_hint": "积极修订：尽可能采纳专家建议，允许新增与删除。",
    },
}

# ─── Logging ───────────────────────────────────────────────────────
_logger = logging.getLogger("tacit_knowledge")
_logger.setLevel(logging.INFO)
if not _logger.handlers:
    _ch = logging.StreamHandler()
    _ch.setFormatter(logging.Formatter(
        "%(asctime)s [%(levelname)s] %(message)s", datefmt="%Y-%m-%d %H:%M:%S"
    ))
    _logger.addHandler(_ch)


def shared_debug_log(level, source, message, data=None):
    """Structured logging."""
    extra = json.dumps(data or {}, ensure_ascii=False, default=str)
    _logger.info("[%s] %s | %s", level, source, f"{message} {extra}")


def shared_agent_debug_log(run_id, level, source, message, data=None):
    """Agent debug logging."""
    extra = json.dumps(data or {}, ensure_ascii=False, default=str)
    _logger.info("[agent:%s][%s] %s | %s", run_id, level, source, f"{message} {extra}")


# ─── Auth Helpers ──────────────────────────────────────────────────
def _mask_api_key(key: str) -> str:
    key = key or ""
    if len(key) > 8:
        return key[:4] + "****" + key[-4:]
    return "****" if key else ""


def _sanitize_model_for_client(model_cfg: dict) -> dict:
    d = {k: v for k, v in model_cfg.items() if k not in ("api_key", "fst_attr_rmrk")}
    d["api_key_masked"] = _mask_api_key(model_cfg.get("api_key", ""))
    d["has_api_key"] = bool((model_cfg.get("api_key") or "").strip())
    return d


# ─── Workbook Context Manager ──────────────────────────────────────
@contextmanager
def _safe_workbook(path, read_only=True, data_only=True):
    """Context manager for openpyxl workbook."""
    wb = None
    try:
        wb = openpyxl.load_workbook(str(path), read_only=read_only, data_only=data_only)
        yield wb
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


# ─── Text Decoding ─────────────────────────────────────────────────
def _decode_text_by_filename(filename: str, raw: bytes) -> str:
    """Decode bytes by file extension, supporting TXT/MD/DOCX/PDF."""
    filename = (filename or "").lower()
    if filename.endswith(".pdf"):
        try:
            from PyPDF2 import PdfReader
            import io
            reader = PdfReader(io.BytesIO(raw))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(pages)
        except Exception:
            return raw.decode("utf-8", errors="replace")
    if filename.endswith(".docx"):
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(raw))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            return "\n".join(paragraphs)
        except Exception:
            return raw.decode("utf-8", errors="replace")
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text_from_file(file_obj) -> str:
    """Read text from uploaded file, supporting TXT, MD, DOCX, PDF."""
    filename = (file_obj.filename or "").lower()
    raw = file_obj.read()
    return _decode_text_by_filename(filename, raw)


def extract_text_from_path(file_path: str) -> str:
    """Read text from a cached workspace file path."""
    p = Path(file_path)
    raw = p.read_bytes()
    return _decode_text_by_filename(p.name, raw)


# ─── Pipeline Helpers ──────────────────────────────────────────────
def _pipeline_prefers_markdown(pipeline_id: str) -> bool:
    if not pipeline_id:
        return False
    with _pipelines_lock:
        pipelines = load_pipelines()
        for p in pipelines:
            if p.get("id") != pipeline_id:
                continue
            sd = p.get("step_data", {}) or {}
            if str(sd.get("step1_output_format", "")).strip().lower() == "markdown":
                return True
            form = sd.get("step1_form_data", {}) or {}
            return str(form.get("output_format", "")).strip().lower() == "markdown"
    return False


def _excel_to_markdown_file(excel_path: str | Path, md_path: str | Path, *, title: str = "") -> None:
    with _safe_workbook(str(excel_path)) as wb:
        lines: list[str] = []
        if title:
            lines.append(f"# {title}")
            lines.append("")
        for ws in wb.worksheets:
            lines.append(f"## 工作表：{ws.title}")
            lines.append("")
            rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
            if not rows:
                lines.append("（空表）")
                lines.append("")
                continue
            max_cols = max((len(r or ()) for r in rows), default=0)
            if max_cols <= 0:
                lines.append("（无可用列）")
                lines.append("")
                continue
            header_row = rows[0] or ()
            headers = [str(h or "").replace("|", "\\|") for h in (list(header_row) + [""] * max_cols)[:max_cols]]
            lines.append("| " + " | ".join(headers) + " |")
            lines.append("| " + " | ".join(["---"] * max_cols) + " |")
            for row in rows[1:]:
                vals = []
                row = row or ()
                for i in range(max_cols):
                    val = row[i] if i < len(row) else ""
                    s = str(val or "").replace("\n", " ").replace("|", "\\|")
                    vals.append(s)
                lines.append("| " + " | ".join(vals) + " |")
            lines.append("")
    Path(md_path).write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def _maybe_generate_markdown_artifact(pipeline_id: str, excel_name: str, *, md_prefix: str, title: str) -> tuple[str, str]:
    if not _pipeline_prefers_markdown(pipeline_id):
        return "", ""
    excel_path = safe_workspace_path(WORKSPACE, excel_name, must_exist=True)
    if not excel_path:
        return "", ""
    stem_id = uuid.uuid4().hex[:8]
    md_name = f"{md_prefix}_{stem_id}.md"
    md_path = WORKSPACE / md_name
    try:
        _excel_to_markdown_file(excel_path, md_path, title=title)
        return md_name, f"/downloads/{md_name}"
    except Exception:
        return "", ""


def get_current_locale(
    *,
    query_lang: str | None = None,
    header_lang: str | None = None,
    pipeline_locale: str | None = None,
    body_locale: str | None = None,
) -> str:
    """Resolve locale from request context.

    Resolution order: query_lang -> pipeline_locale -> body_locale -> header_lang -> DEFAULT_LANG.
    """
    return resolve_locale(
        query_lang=query_lang,
        header_lang=header_lang,
        pipeline_locale=pipeline_locale or body_locale,
    )


# ─── JSON Parsing (6-layer fallback) ───────────────────────────────
def _extract_balanced_json_slice(text: str, open_ch: str, close_ch: str) -> str:
    start = text.find(open_ch)
    if start < 0:
        return ""
    depth = 0
    in_str = False
    esc = False
    for i in range(start, len(text)):
        ch = text[i]
        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
            continue
        if ch == '"':
            in_str = True
            continue
        if ch == open_ch:
            depth += 1
        elif ch == close_ch:
            depth -= 1
            if depth == 0:
                return text[start : i + 1]
    return ""


def _repair_json_text(text: str) -> str:
    if not text:
        return text
    t = text.strip()
    t = t.replace("“", '"').replace("”", '"').replace("‘", "'").replace("’", "'")
    t = re.sub(r",\s*([}\]])", r"\1", t)
    return t


def _extract_json_from_text(text: str) -> str:
    text = (text or "").strip()
    if not text:
        return ""
    for marker in ["```json", "```"]:
        if marker in text:
            parts = text.split(marker, 1)
            if len(parts) > 1:
                inner = parts[1].split("```", 1)[0].strip()
                if inner:
                    text = inner
                    break
    text = text.strip()
    if text.startswith("["):
        slice_text = _extract_balanced_json_slice(text, "[", "]")
        return slice_text or text
    if text.startswith("{"):
        slice_text = _extract_balanced_json_slice(text, "{", "}")
        return slice_text or text
    for open_ch, close_ch in (("[", "]"), ("{", "}")):
        slice_text = _extract_balanced_json_slice(text, open_ch, close_ch)
        if slice_text:
            return slice_text
    return text


def _items_from_parsed_root(parsed) -> list | None:
    if isinstance(parsed, list):
        return parsed
    if not isinstance(parsed, dict):
        return None
    for key in (
        "items", "data", "results", "result", "entries", "knowledge",
        "knowledge_items", "extracted", "extracted_items", "知识", "知识条目", "条目",
    ):
        val = parsed.get(key)
        if isinstance(val, list):
            return val
    return None


def _extract_top_level_json_objects(text: str) -> list[str]:
    objects = []
    depth = 0
    in_str = False
    esc = False
    start = -1
    for i, ch in enumerate(text):
        if in_str:
            if esc: esc = False
            elif ch == "\\": esc = True
            elif ch == '"': in_str = False
            continue
        if ch == '"': in_str = True; continue
        if ch == "{":
            if depth == 0: start = i
            depth += 1; continue
        if ch == "}":
            if depth > 0:
                depth -= 1
                if depth == 0 and start >= 0:
                    objects.append(text[start:i + 1])
                    start = -1
    return objects


def _parse_extracted_items(raw_text: str) -> tuple[list, str]:
    """Parse LLM extraction output with 6-layer robust fallback."""
    cleaned = _extract_json_from_text(raw_text or "")
    if not cleaned.strip():
        return [], "empty"
    for candidate in (cleaned, _repair_json_text(cleaned)):
        if not candidate.strip():
            continue
        try:
            parsed = json.loads(candidate)
            items = _items_from_parsed_root(parsed)
            if items is not None:
                return items, "json_list" if isinstance(parsed, list) else "json_items"
        except Exception:
            pass
    # JSONL fallback
    line_items = []
    for line in cleaned.splitlines():
        line = line.strip().rstrip(",")
        if not line.startswith("{"):
            continue
        try:
            obj = json.loads(_repair_json_text(line))
            if isinstance(obj, dict):
                line_items.append(obj)
        except Exception:
            continue
    if line_items:
        return line_items, "jsonl"
    # Object recovery
    recovered = []
    for obj_text in _extract_top_level_json_objects(_repair_json_text(cleaned)):
        try:
            obj = json.loads(obj_text)
            if isinstance(obj, dict):
                recovered.append(obj)
        except Exception:
            continue
    if recovered:
        return recovered, "object_recovery"
    # Last resort: scan raw text
    for obj_text in _extract_top_level_json_objects(raw_text or ""):
        try:
            obj = json.loads(_repair_json_text(obj_text))
            if isinstance(obj, dict):
                recovered.append(obj)
        except Exception:
            continue
    if recovered:
        return recovered, "object_recovery"
    return [], "empty"


# ─── LLM Config Management ─────────────────────────────────────────
def load_preset_overrides():
    if not PRESET_OVERRIDES_PATH.exists():
        return {}
    try:
        with open(str(PRESET_OVERRIDES_PATH), "r", encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, IOError):
        return {}


def save_preset_overrides(overrides):
    with _models_lock:
        PRESET_OVERRIDES_PATH.write_text(json.dumps(overrides, ensure_ascii=False, indent=2), encoding="utf-8")


def load_base_presets():
    if not LLM_CONFIG_PATH.exists():
        return []
    with open(str(LLM_CONFIG_PATH), "r", encoding="utf-8") as f:
        cfg = yaml.safe_load(f) or {}
    presets = [dict(m) for m in cfg.get("presets", [])]
    if LLM_CONFIG_LOCAL_PATH.exists():
        with open(str(LLM_CONFIG_LOCAL_PATH), "r", encoding="utf-8") as f:
            local_cfg = yaml.safe_load(f) or {}
        local_by_name = local_cfg.get("presets") or {}
        if isinstance(local_by_name, dict):
            for entry in presets:
                name = entry.get("name")
                if name and name in local_by_name:
                    override = local_by_name[name]
                    if isinstance(override, dict):
                        entry.update(override)
    return presets


def load_llm_config():
    models = []
    overrides = load_preset_overrides()
    for m in load_base_presets():
        entry = dict(m)
        if entry.get("name") in overrides:
            entry.update(overrides[entry["name"]])
        if entry.get("url"):
            api_type = (entry.get("api_type") or API_TYPE_OPENAI).strip().lower()
            entry["url"] = normalize_llm_url(entry["url"], api_type)
        entry["is_preset"] = True
        models.append(entry)
    custom = []
    if CUSTOM_MODELS_PATH.exists():
        try:
            with open(str(CUSTOM_MODELS_PATH), "r", encoding="utf-8") as f:
                custom = json.load(f)
        except (json.JSONDecodeError, IOError):
            custom = []
    for m in custom:
        entry = dict(m)
        if entry.get("url"):
            api_type = (entry.get("api_type") or API_TYPE_OPENAI).strip().lower()
            entry["url"] = normalize_llm_url(entry["url"], api_type)
        entry["is_preset"] = False
        models.append(entry)
    return models


def save_custom_models(custom_list):
    clean = []
    for m in custom_list:
        d = {k: v for k, v in m.items() if k != "is_preset"}
        clean.append(d)
    with _models_lock:
        CUSTOM_MODELS_PATH.write_text(json.dumps(clean, ensure_ascii=False, indent=2), encoding="utf-8")


def get_model_by_name(name):
    n = (name or "").strip().lower()
    for m in load_llm_config():
        if (m.get("name") or "").strip().lower() == n:
            return m
    return None


# ─── Pipeline Persistence ──────────────────────────────────────────
def load_pipelines():
    if not PIPELINES_PATH.exists():
        return []
    try:
        with open(str(PIPELINES_PATH), "r", encoding="utf-8") as f:
            return json.load(f)
    except (json.JSONDecodeError, IOError):
        return []


def save_pipelines(pipelines):
    with _pipelines_lock:
        PIPELINES_PATH.write_text(json.dumps(pipelines, ensure_ascii=False, indent=2), encoding="utf-8")


# ─── File Save/Upload ──────────────────────────────────────────────
def save_upload(file_obj, prefix="upload"):
    ext = os.path.splitext(file_obj.filename or ".bin")[1] or ".bin"
    name = f"{prefix}_{uuid.uuid4().hex[:8]}{ext}"
    path = WORKSPACE / name
    file_obj.save(str(path))
    return str(path)


def run_script(script_name, args_list):
    script_path = SCRIPT_DIR / "scripts" / script_name
    if not script_path.exists():
        raise FileNotFoundError(f"Script not found: {script_name}")
    args = [sys.executable, str(script_path)] + args_list
    result = subprocess.run(args, capture_output=True, text=True, timeout=120, cwd=str(SCRIPT_DIR))
    if result.returncode != 0:
        raise RuntimeError(f"Script {script_name} failed: {result.stderr[:500]}")
    try:
        return json.loads(result.stdout)
    except json.JSONDecodeError:
        return {"output": result.stdout}


# ─── Excel Helpers ─────────────────────────────────────────────────
def _trim_excel_rows(rows, max_rows=200, max_cols=40):
    trimmed = []
    for r in rows[:max_rows]:
        trimmed.append(list(r[:max_cols]) if len(r) > max_cols else list(r))
    return trimmed


# ─── Extract Helpers ───────────────────────────────────────────────
def _align_item_keys_to_template(item: dict, target_columns: list[str]) -> dict:
    if not target_columns:
        return item
    row = {k: item.get(k, "") for k in target_columns}
    return row


def _normalize_extracted_items(items: list, target_columns: list[str]) -> list:
    if not target_columns:
        return items
    return [_align_item_keys_to_template(it, target_columns) for it in items]


def _pick_text(item: dict, keys: tuple[str, ...]) -> str:
    for k in keys:
        v = item.get(k, "")
        if v and str(v).strip():
            return str(v).strip()
    return ""


def _extract_item_content(item: dict, target_columns: list | None = None) -> str:
    priority = ("知识描述", "判断逻辑", "适用条件", "具体方法", "可执行知识",
                 "操作要点", "专家经验", "隐性信号", "经验判断")
    keys = tuple(priority)
    if target_columns:
        keys = tuple([c for c in target_columns if c not in
                      ("置信度", "知识分类", "来源", "来源文档", "来源位置",
                       "知识类型", "知识引用", "贡献专家", "确认专家", "证据数", "突破数",
                       "适用边界", "例外情形", "访谈方向", "备注", "反模式/踩坑提示",
                       "隐性信号", "信号类型")])
    return _pick_text(item, keys)


def _extract_item_confidence_rank(item: dict) -> int:
    conf = str(item.get("置信度", "")).strip()
    return {"高": 3, "中": 2, "低": 1, "极高": 4, "极低": 0}.get(conf, 2)


def _extract_item_richness(item: dict) -> int:
    score = 0
    fields = ("知识描述", "判断逻辑", "适用条件", "反模式/踩坑提示",
              "经验判断", "适用边界", "例外情形", "来源文档", "来源位置",
              "贡献专家", "证据数", "突破数")
    for f in fields:
        v = item.get(f, "")
        if v and str(v).strip() and str(v).strip() not in ("无", "N/A", "暂无"):
            score += 1
    return score


# ─── Revision Helpers ──────────────────────────────────────────────
def _normalize_revision_style(style: str) -> str:
    style = (style or "").strip()
    return style if style in REVISION_STYLE_RULES else "标准修订"


def _normalize_extract_style(style: str) -> str:
    style = (style or "").strip()
    return style if style in EXTRACT_STYLE_RULES else "标准萃取"


def _normalize_expert_text_for_align(expert_text: str) -> str:
    return (expert_text or "").strip()


def _text_is_explicit_no_opinion(expert_text: str) -> bool:
    t = (expert_text or "").strip()
    if not t:
        return False
    markers = ["无意见", "没有意见", "没问题", "无需修改", "无修订", "不用改",
               "没有修改意见", "no opinion", "no changes", "no revision", "无需调整"]
    return any(m in t.lower() for m in markers)


def _alignment_llm_guard_rules() -> str:
    return (
        "\n\n【LLM 行为边界 — 必须严格遵守】\n"
        "1. 只输出 JSON 数组；无修订时输出 []。\n"
        "2. 每个修订必须能在专家意见中找到明确的依据——不能凭空编造。\n"
        "3. 禁止根据自己对知识稿的理解主动「优化」或「改写」条目。\n"
        "4. 若专家意见与知识稿内容无关，不要强行关联。\n"
        "5. old_value 必须与当前表中内容完全一致，不得猜测或改写。\n"
    )


def _to_int(val, default=0) -> int:
    try:
        return int(float(str(val).strip()))
    except (ValueError, TypeError):
        return default


def _normalize_revision_action(action: str) -> str:
    action = (action or "").strip().lower()
    if action in ("modify", "supplement", "add", "delete"):
        return action
    if action in ("修改", "修正", "更新"):
        return "modify"
    if action in ("补充", "追加"):
        return "supplement"
    if action in ("新增", "添加", "加入"):
        return "add"
    if action in ("删除", "移除"):
        return "delete"
    return "modify"


# Alias for backward compatibility
_old_import_fix = is_step2_preextract_filename
