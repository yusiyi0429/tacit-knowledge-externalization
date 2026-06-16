#!/usr/bin/env python3
"""对公普惠客户潜力营销结构化萃取流水线编排脚本。

运行前请确保 backend server 已启动：
    cd backend && python app_server.py --host 127.0.0.1 --port 5000

脚本流程：
    1. 检查服务可用性
    2. 向验证知识库注入本场景测试用例（source=test_phmarketing_structured）
    3. 向案例知识库注入 Step5 回放所需案例
    4. 创建 pipeline 并手动回填 Step1/Step2 产物
    5. 顺序执行 Step3 对齐 → Step4 编译 → Step5 回放
    6. 清理带 source 标记的验证知识库用例
"""
from __future__ import annotations

import json
import logging
import os
import shutil
import sqlite3
import sys
import uuid
from pathlib import Path

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipeline_artifacts import workspace_path_for  # noqa: E402

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
BASE = "http://127.0.0.1:5000"
WORKSPACE = PROJECT_ROOT / "data" / "workspace"
SOURCE_XLSX = (
    PROJECT_ROOT
    / "data"
    / "samples"
    / "step2-文档萃取"
    / "对公普惠客户潜力营销_结构化萃取模板.xlsx"
)
MODEL = "DeepSeek-V4-Flash"
DOMAIN = "对公普惠客户潜力营销"
SCENARIO = "对公普惠客户潜力营销"
KB_SOURCE = "test_phmarketing_structured"

KB_CASES = [
    {
        "description": "省级专精特新企业，成立3年，账户状态正常，金融总量800万、存款日均高，无申贷记录，未办理我行普惠信用快贷。",
        "expert_conclusion": "拒绝。客户虽满足基本准入条件，但金融总量高、存款日均高且无申贷记录，属无融资需求客户；建议推荐现金池、一户通、票据池、大额存单、基金等结算/投资类产品。",
        "facts": {
            "客户归属机构": "当前网点",
            "账户状态": "正常",
            "科技型企业标签": "对公客户是否是专精特新企业",
            "授信客户标志": "0",
            "成立年限": "3年",
            "实控人年龄": "45",
            "实控人身份": "中国大陆居民",
            "征信记录": "良好",
            "他行信贷合作": "正常",
            "金融总量余额": "800万",
            "活期存款余额": "50万",
            "定期存款年日均": "70万",
            "申贷次数": "0",
        },
        "difficulty": "normal",
        "tags": "对公普惠,科技型企业,无融资需求",
        "source": KB_SOURCE,
    },
    {
        "description": "国家高新技术企业，成立1年，账户状态正常，未办理信用快贷，实控人年龄32岁。",
        "expert_conclusion": "拒绝。客户成立年限仅1年，不满足信用快贷基本准入条件；建议持续跟踪，待成立年限达标后重新评估，期间做好结算服务。",
        "facts": {
            "客户归属机构": "当前网点",
            "账户状态": "正常",
            "科技型企业标签": "对公客户国家创新型企业标志",
            "授信客户标志": "0",
            "成立年限": "1年",
            "实控人年龄": "32",
            "实控人身份": "中国大陆居民",
            "征信记录": "良好",
            "他行信贷合作": "正常",
        },
        "difficulty": "edge",
        "tags": "对公普惠,科技型企业,不满足准入",
        "source": KB_SOURCE,
    },
    {
        "description": "科技型中小企业，成立4年，账户状态正常，有真实融资需求但预测算额度低于预期，曾申贷2次。",
        "expert_conclusion": "条件通过。客户符合基本准入条件，但预测算额度较低、产品匹配度不足；建议推荐抵押快贷/善营贷提升额度，或善新贷/政策性贴息产品降低利率。",
        "facts": {
            "客户归属机构": "当前网点",
            "账户状态": "正常",
            "科技型企业标签": "对公客户科技型中小企业标志",
            "授信客户标志": "0",
            "成立年限": "4年",
            "实控人年龄": "42",
            "实控人身份": "中国大陆居民",
            "征信记录": "良好",
            "他行信贷合作": "正常",
            "申贷次数": "2",
            "预测算额度": "30万",
        },
        "difficulty": "hard",
        "tags": "对公普惠,科技型企业,产品匹配度不足",
        "source": KB_SOURCE,
    },
]


def _create_session() -> requests.Session:
    """创建带简单重试的 HTTP Session（POST 也允许重试）。"""
    session = requests.Session()
    retry_strategy = Retry(
        total=2,
        backoff_factor=1,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=frozenset(["HEAD", "GET", "POST", "PUT", "DELETE", "OPTIONS"]),
    )
    adapter = HTTPAdapter(max_retries=retry_strategy)
    session.mount("http://", adapter)
    session.mount("https://", adapter)
    return session


SESSION = _create_session()


def _require_status_ok(data: dict, context: str) -> None:
    """校验 API 返回 data['status'] == 'ok'，否则抛出 RuntimeError。"""
    if data.get("status") != "ok":
        error = data.get("error", data)
        raise RuntimeError(f"{context} failed: {error}")


def _require_file(path: Path, context: str) -> None:
    """校验文件存在，否则抛出 FileNotFoundError。"""
    if not path.is_file():
        raise FileNotFoundError(f"{context}: required file not found: {path}")


def check_server_ready() -> None:
    """检查后端服务是否就绪；未就绪时打印错误并退出进程。"""
    try:
        r = SESSION.get(BASE + "/api/pipelines", timeout=10)
        r.raise_for_status()
    except requests.RequestException as exc:
        logger.error("Backend server is not ready at %s: %s", BASE, exc)
        logger.error("Please start it with: cd backend && python app_server.py --host 127.0.0.1 --port 5000")
        sys.exit(1)


def create_pipeline() -> str:
    """创建 PHM-structured pipeline，返回 pipeline ID。"""
    payload = {
        "name": "PHM-structured",
        "scenario": SCENARIO,
        "domain": DOMAIN,
    }
    r = SESSION.post(BASE + "/api/pipelines", json=payload, timeout=30)
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "create_pipeline")
    pid = data["pipeline"]["id"]
    logger.info("Created pipeline %s", pid)
    return pid


def seed_step2_file(pipeline_id: str) -> str:
    """把结构化 Excel 复制到 workspace 并注册为 Step1/Step2 产物，同时推进 step_status。"""
    _require_file(SOURCE_XLSX, "seed_step2_file")

    # 复制 Step1 占位文件，确保 step1_output_file 引用真实存在
    step1_placeholder = workspace_path_for(WORKSPACE, pipeline_id, "step1", "template_uploaded.xlsx")
    step1_placeholder.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SOURCE_XLSX, step1_placeholder)

    fname = f"preextract_{uuid.uuid4().hex[:8]}.xlsx"
    dest = workspace_path_for(WORKSPACE, pipeline_id, "step2", fname)
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SOURCE_XLSX, dest)

    # 动态计算萃取条数（表头占 1 行）
    from openpyxl import load_workbook

    wb = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    ws = wb.active
    extracted_count = ws.max_row - 1
    wb.close()

    r = SESSION.get(BASE + f"/api/pipelines/{pipeline_id}", timeout=30)
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "seed_step2_file/get_pipeline")
    p = data["pipeline"]
    sd = p.get("step_data", {})

    sd["step1_form_data"] = {
        "scenario_name": "对公普惠客户潜力营销",
        "scenario_content": (
            "基于网点存量对公客户，运用行内外科技型企业标签与风险数据，识别具备科技资质但未办理普惠信用快贷的潜在客户。"
            "通过“客户筛选→客户数据匹配→原因归因→决策建议”的逐层逻辑，输出分层营销建议与产品推荐。"
        ),
        "sub_scenarios": [
            {
                "name": "科技型企业普惠贷款营销",
                "content": (
                    "指建设银行拟运用互联网和移动通信等信息通信技术，基于风险数据和风险模型进行交叉验证和风险管理，"
                    "线上自动受理贷款申请及开展风险评估，并完成授信审批、合同签订、贷款支付、贷后管理等核心业务环节操作，"
                    "向符合条件的小微企业法人或非法人组织发放的，用于借款人日常经营周转的信用类流动资金贷款业务。"
                ),
            }
        ],
    }
    sd["step1_output_file"] = "template_uploaded.xlsx"
    sd["step1_download_url"] = "/downloads/template_uploaded.xlsx"
    sd["step1_output_format"] = "excel"
    sd["step1_knowledge_columns"] = [
        "场景", "场景说明", "子场景", "子场景说明", "步骤",
        "具体方法", "知识引用", "规则引用", "专业术语",
        "关键输出-名称", "关键输出-描述",
    ]
    sd["step1_user_knowledge_columns"] = sd["step1_knowledge_columns"]
    sd["step1_columns_enriched"] = False
    sd["step2_output_file"] = fname
    sd["step2_download_url"] = f"/downloads/{fname}"
    sd["step2_extracted_count"] = extracted_count
    sd["step2_source_count"] = 1
    sd["step2_dedup_count"] = 0
    sd["step2_extract_style"] = "标准萃取"

    # 标记 Step1/Step2 已完成，并把 current_step 推进到 2
    step_status = p.get("step_status", {})
    step_status["1"] = "done"
    step_status["2"] = "done"

    r = SESSION.put(
        BASE + f"/api/pipelines/{pipeline_id}",
        json={
            "step_data": sd,
            "step_status": step_status,
            "current_step": 2,
        },
        timeout=30,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "seed_step2_file/put_pipeline")
    logger.info("Seeded step2 file %s (extracted_count=%d)", fname, extracted_count)
    return fname


def ensure_verification_cases() -> None:
    """若验证知识库中无本场景用例，则写入带 source 标记的验证用例。"""
    r = SESSION.get(
        BASE + "/api/kb/verification_cases",
        params={"source": KB_SOURCE, "limit": 100},
        timeout=30,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "ensure_verification_cases/list")
    existing = data.get("cases", [])
    if existing:
        logger.info("KB verification_cases already has %d case(s) for source=%s", len(existing), KB_SOURCE)
        return

    created = 0
    for i, c in enumerate(KB_CASES, start=1):
        payload = {
            "name": f"{SCENARIO}-structured-{i}",
            "description": c["description"],
            "input": c["facts"],
            "expected_output": {"conclusion": c["expert_conclusion"]},
            "tags": c["tags"],
            "source": c["source"],
            "skill_id": "",
        }
        r = SESSION.post(
            BASE + "/api/kb/verification_cases",
            json=payload,
            timeout=30,
        )
        r.raise_for_status()
        data = r.json()
        _require_status_ok(data, f"ensure_verification_cases/create/{i}")
        created += 1
    logger.info("Seeded %d KB verification_cases", created)


def ensure_step5_cases() -> None:
    """向案例知识库注入 Step5 回放所需的案例（case_source=kb 读取的来源）。"""
    r = SESSION.get(
        BASE + "/api/kb/cases",
        params={"domain": DOMAIN, "scenario": SCENARIO, "limit": 10},
        timeout=30,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "ensure_step5_cases/list")
    existing = data.get("cases", [])
    if existing:
        logger.info("KB cases already has %d case(s) for domain=%s scenario=%s", len(existing), DOMAIN, SCENARIO)
        return

    cases = [{**c, "domain": DOMAIN, "scenario": SCENARIO} for c in KB_CASES]
    r = SESSION.post(
        BASE + "/api/kb/cases",
        json={"cases": cases},
        timeout=30,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "ensure_step5_cases/create")
    logger.info("Seeded %d KB cases", data.get("created", 0))


def run_step3_finalize(pipeline_id: str) -> None:
    """执行 Step3 知识对齐 / 终版生成。"""
    r = SESSION.post(
        BASE + "/api/step3/finalize",
        data={"pipeline_id": pipeline_id, "model": MODEL, "style": "标准修订"},
        timeout=300,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "run_step3_finalize")
    logger.info(
        "Step3 finalized: output_file=%s, aligned_file=%s",
        data.get("output_file"),
        data.get("aligned_file"),
    )


def run_step4_compile(pipeline_id: str) -> None:
    """执行 Step4 智能转化 / SKILL 编译。"""
    r = SESSION.post(
        BASE + "/api/step4/compile",
        data={"pipeline_id": pipeline_id, "model": MODEL},
        timeout=300,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "run_step4_compile")
    artifacts = data.get("artifacts_download") or {}
    logger.info(
        "Step4 compiled: input_kind=%s, skill=%s, artifacts=%s",
        data.get("input_kind"),
        data.get("download_name"),
        json.dumps(list(artifacts.keys()), ensure_ascii=False),
    )


def run_step5_replay(pipeline_id: str) -> dict:
    """执行 Step5 验证回放，返回接口完整响应数据。"""
    r = SESSION.post(
        BASE + "/api/step5/replay",
        data={
            "pipeline_id": pipeline_id,
            "judge_model": MODEL,
            "case_source": "kb",
            "kb_domain": DOMAIN,
            "kb_scenario": SCENARIO,
            "kb_limit": 10,
        },
        timeout=300,
    )
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "run_step5_replay")
    logger.info(
        "Step5 replay: run_id=%s, hit_rate=%s",
        data.get("run_id"),
        data.get("hit_rate"),
    )
    return data


def _kb_db_path() -> Path:
    """返回知识库 SQLite 文件路径（优先 KB_DB_PATH 环境变量）。"""
    env_path = os.environ.get("KB_DB_PATH")
    if env_path:
        return Path(env_path)
    return PROJECT_ROOT / "data" / "kb" / "knowledge_base.db"


def cleanup_kb_cases_table() -> None:
    """直接清理 kb_cases 表中 source=test_phmarketing_structured 的记录；失败仅记录警告。"""
    db_path = _kb_db_path()
    if not db_path.is_file():
        logger.warning("KB database not found at %s, skipping kb_cases cleanup", db_path)
        return
    try:
        conn = sqlite3.connect(str(db_path))
        try:
            cur = conn.execute("DELETE FROM kb_cases WHERE source = ?", (KB_SOURCE,))
            conn.commit()
            logger.info("Cleaned %d row(s) from kb_cases where source=%s", cur.rowcount, KB_SOURCE)
        finally:
            conn.close()
    except sqlite3.Error as exc:
        logger.warning("Exception cleaning kb_cases table: %s", exc)


def cleanup_verification_cases() -> None:
    """删除 source=test_phmarketing_structured 的验证知识库用例；失败仅记录警告。"""
    try:
        r = SESSION.get(
            BASE + "/api/kb/verification_cases",
            params={"source": KB_SOURCE, "limit": 100},
            timeout=30,
        )
        r.raise_for_status()
        data = r.json()
        if data.get("status") != "ok":
            logger.warning("Failed to list verification_cases for cleanup: %s", data.get("error"))
            return
        cases = data.get("cases", [])
        if not cases:
            logger.info("No verification_cases with source=%s to clean up", KB_SOURCE)
            return
        for c in cases:
            case_uid = c.get("case_uid")
            if not case_uid:
                continue
            try:
                dr = SESSION.delete(BASE + f"/api/kb/verification_cases/{case_uid}", timeout=30)
                dr.raise_for_status()
                ddata = dr.json()
                if ddata.get("status") == "ok":
                    logger.info("Deleted verification_case %s", case_uid)
                else:
                    logger.warning("Failed to delete verification_case %s: %s", case_uid, ddata.get("error"))
            except requests.RequestException as exc:
                logger.warning("Exception deleting verification_case %s: %s", case_uid, exc)
    except requests.RequestException as exc:
        logger.warning("Exception listing verification_cases for cleanup: %s", exc)


def verify_final_pipeline(pipeline_id: str) -> dict:
    """获取最终 pipeline 状态并校验所有步骤已完成。"""
    r = SESSION.get(BASE + f"/api/pipelines/{pipeline_id}", timeout=30)
    r.raise_for_status()
    data = r.json()
    _require_status_ok(data, "verify_final_pipeline")
    p = data["pipeline"]
    step_status = p.get("step_status", {})
    current_step = p.get("current_step")
    for step in ("1", "2", "3", "4", "5"):
        if step_status.get(step) != "done":
            raise RuntimeError(f"Step {step} is not done: {step_status}")
    if current_step != 5:
        raise RuntimeError(f"Expected current_step=5, got {current_step}")
    logger.info("Final pipeline state OK: current_step=%s, step_status=%s", current_step, step_status)
    return p


def main() -> None:
    """编排完整流水线并负责最终清理。"""
    check_server_ready()
    ensure_verification_cases()
    ensure_step5_cases()
    pid = create_pipeline()
    try:
        seed_step2_file(pid)
        run_step3_finalize(pid)
        run_step4_compile(pid)
        step5_result = run_step5_replay(pid)
        verify_final_pipeline(pid)
        hit_rate = step5_result.get("hit_rate")
        logger.info("\nPipeline ID: %s", pid)
        logger.info("Hit rate: %s", hit_rate)
    finally:
        cleanup_verification_cases()
        cleanup_kb_cases_table()


if __name__ == "__main__":
    main()
